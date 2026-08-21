import sys
from pathlib import Path
from docx import Document

HERE = Path(__file__).resolve().parent
doc = Document(str(HERE / "Report chính.docx"))

lo, hi = int(sys.argv[1]), int(sys.argv[2])
for i in range(lo, min(hi + 1, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"[P{i} {p.style.name}] {p.text}")

for ti, table in enumerate(doc.tables):
    print(f"\n--- TABLE {ti} ---")
    for row in table.rows:
        print("  | " + " | ".join(c.text for c in row.cells))
