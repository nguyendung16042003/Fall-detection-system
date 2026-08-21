from pathlib import Path
from docx import Document

HERE = Path(__file__).resolve().parent
doc = Document(str(HERE / "Report chính.docx"))

print(f"So doan van: {len(doc.paragraphs)}, so bang: {len(doc.tables)}")
print("\n=== MUC LUC (cac dong dung style Heading) ===")
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else ""
    if style.startswith("Heading") or style.startswith("Title"):
        print(f"[P{i:5d} {style:12s}] {p.text}")
