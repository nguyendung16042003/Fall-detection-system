import sys
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
prs = Presentation(str(HERE / "Slide chính.pptx"))
print(f"Tong so slide: {len(prs.slides)}")
print(f"Kich thuoc slide: {prs.slide_width} x {prs.slide_height} EMU")

target = set(int(x) for x in sys.argv[1:]) if len(sys.argv) > 1 else None

def dump_shape(shape, depth=0):
    pad = "  " * depth
    if shape.has_text_frame and shape.text_frame.text.strip():
        print(f"{pad}--- [TEXT name={shape.name!r}] ---")
        for para in shape.text_frame.paragraphs:
            line = "".join(r.text for r in para.runs)
            if line.strip():
                print(f"{pad}  {line}")
    elif shape.has_table:
        print(f"{pad}--- [TABLE name={shape.name!r}] ---")
        for row in shape.table.rows:
            print(f"{pad}  | " + " | ".join(c.text for c in row.cells))
    elif shape.shape_type == 6:  # GROUP
        print(f"{pad}--- [GROUP name={shape.name!r}] ---")
        for sub in shape.shapes:
            dump_shape(sub, depth + 1)
    elif shape.shape_type == 13:  # PICTURE
        print(f"{pad}--- [PICTURE name={shape.name!r}] ---")
    # bo qua freeform/autoshape khong text (trang tri)


for i, slide in enumerate(prs.slides, start=1):
    if target and i not in target:
        continue
    print(f"\n{'='*30} SLIDE {i} (layout: {slide.slide_layout.name}) {'='*30}")
    for shape in slide.shapes:
        dump_shape(shape)
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    if notes.strip():
        print(f"--- [SPEAKER NOTES] ---\n  {notes}")
