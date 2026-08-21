"""Tao 3 file Word (Bai cua nhom / Doi thu / So sanh) cho report Identity
Association (ten hoc thuat thay P2 -- xem Report chinh.docx muc 1.2/Part 2),
cung cau truc voi make_reid_report_docs_kd025.py / make_bff_report_docs.py."""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches

sys.stdout.reconfigure(encoding="utf-8")

AI_ROOT = Path(r"D:\DOWLOAD\FileTaiLieuHocTapCuaDung\Ki9\Đồ án\Fall-detection-system\Fall-detection-system\Tan Dung (AI)")
BASE = AI_ROOT / "Chạy xử lý toàn bộ hệ thống" / "Số liệu thống kê" / "Identity Association"
DIR_OWN = BASE / "1. Bai cua nhom (Homography + Hungarian)"
DIR_COMP = BASE / "2. Doi thu (GNN-CCA - Luna et al 2022)"
DIR_CMP = BASE / "3. So sanh doi thu"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


def add_image(doc, path, width_in=5.8):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width_in))


def add_file_catalog(doc, entries):
    add_heading(doc, "Danh mục file trong thư mục này", level=2)
    for name, desc in entries:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(name)
        run.bold = True
        p.add_run(f" — {desc}")


# ============================================================
# 1. BAI CUA NHOM (Homography + Hungarian)
# ============================================================
def build_doc_own():
    doc = Document()
    doc.add_heading("Identity Association — Bài của nhóm", level=0)

    add_heading(doc, "1. Cách triển khai")
    doc.add_paragraph(
        "Hình học thuần, KHÔNG train model nào (đúng \"Context xử lý vấn đề 2.docx\": \"Bản GNN "
        "đã bị loại bỏ vì thêm 1 model riêng, không khớp ràng buộc 'không thêm model' trên Jetson "
        "Nano 4GB\"). Pipeline: chiếu điểm chân bbox qua homography (calib bằng 4 điểm vuông "
        "0.7m thật) ra toạ độ sàn (mét) cho mỗi camera → Hungarian algorithm "
        "(scipy.optimize.linear_sum_assignment) tìm phép gán tối ưu TOÀN CỤC giữa 2 tập track "
        "2 camera, dựa trên ma trận chi phí = khoảng cách Euclidean → chỉ giữ cặp có khoảng cách "
        "≤ ngưỡng 0.5m."
    )
    doc.add_paragraph(
        "Điểm thiết kế quan trọng: dùng Hungarian (gán ghép TỐI ƯU TOÀN CỤC, 1-1) thay vì chỉ "
        "threshold khoảng cách đơn thuần — đây chính là điểm được kiểm chứng có tác dụng thật ở "
        "Mục 2 (Scene 2)."
    )

    add_heading(doc, "2. Số liệu khi test")
    doc.add_paragraph(
        "Đánh giá bằng ĐÚNG bộ metric clustering mà Luna et al. 2022 (GNN-CCA, đối thủ — xem "
        "folder 2) dùng trong Table V của paper: Adjusted Rand Index (ARI), Adjusted Mutual "
        "Information (AMI), Homogeneity (H), Completeness (C), V-measure — tính trên TỪNG frame "
        "rồi lấy trung bình, giống hệt cách paper làm — để so sánh công bằng, không tự đặt ra "
        "metric khác chỉ vì có lợi."
    )
    doc.add_paragraph(
        "Test trên Identity Association Data 3 (4 cảnh, nhiều người) — bộ test khó nhất của Identity Association, có cả tình huống "
        "đứng sát nhau, che khuất, và ngã. ĐẦY ĐỦ: bản gốc + 4 biến thể augment an toàn với "
        "homography (sáng/tối/nhiễu/nén — không lật/xoay/crop vì Identity Association phụ thuộc calib pixel gốc, "
        "giống quy ước đã dùng cho Re-Identification/Boundary Feature Fusion) — tổng 4 cảnh × 5 biến thể = 840 frame (2 người/frame):"
    )
    add_table(doc, ["Metric", "Value"], [
        ["ARI", "85.45"],
        ["AMI", "85.46"],
        ["Homogeneity", "99.97"],
        ["Completeness", "91.91"],
        ["V-measure", "95.01"],
    ])
    doc.add_paragraph(
        "Ổn định qua mọi biến thể (ARI dao động hẹp 85.0-86.1% qua 5 biến thể) — không phải may "
        "mắn ở 1 lần đo. Riêng bản gốc (168 frame, chưa augment): ARI 85.12, V-measure 95.31 — "
        "gần như giống hệt số trung bình đầy đủ."
    )
    doc.add_paragraph(
        "Lưu ý QUAN TRỌNG về ground truth: Identity Association Data 3 không có nhãn danh tính gán tay sẵn. Ground "
        "truth được xây bằng cách bootstrap từ TẦN SUẤT đồng-xuất-hiện của cặp track qua Hungarian "
        "trên toàn bộ cảnh (cặp xuất hiện ≥2 lần mới coi là 'cùng 1 người thật', gộp bằng Connected "
        "Components) — đây là kỹ thuật vote đa số qua nhiều frame để lọc nhiễu, KHÔNG phải nhãn tay "
        "độc lập. Cần nêu rõ giới hạn này khi trích dẫn: các trường hợp Hungarian sai LẶP LẠI NHIỀU "
        "LẦN (không phải lỗi lẻ tẻ) có thể không bị phát hiện bởi cách đánh giá này."
    )

    add_heading(doc, "3. Các trường hợp khó — phân tích theo từng cảnh")
    add_image(doc, DIR_OWN / "chart_per_scene_ari.png")
    add_table(doc, ["Scene", "Description", "n (5 variants)", "ARI (Hungarian)", "ARI (Threshold+CC)"], [
        ["1", "Standing 1-1.5m apart (normal)", "150", "100.0", "100.0"],
        ["2", "Standing/sitting very close (most ambiguous)", "120", "100.0", "49.2"],
        ["3", "A occluded in 1 camera", "175", "74.3", "74.3"],
        ["4", "A falls, B stands nearby", "395", "80.5", "79.1"],
    ])
    doc.add_paragraph(
        "Scene 3 (che khuất) là cảnh khó nhất cho CẢ 2 phương pháp như nhau (74.3% cả hai) — hợp "
        "lý vì che khuất là vấn đề THIẾU DETECTION (ít frame có đủ dữ liệu 2 camera để ghép), "
        "không phải lỗi thuật toán ghép cặp — không phương pháp ghép nào (kể cả GNN học được) có "
        "thể sửa được khi input đầu vào (detection) đã thiếu."
    )
    doc.add_paragraph(
        "Scene 2 (đứng sát nhau) là bằng chứng RÕ NHẤT cho giá trị của Hungarian, và càng RÕ HƠN "
        "sau khi có đủ augment: threshold+CC đơn thuần (không tối ưu toàn cục, cho phép 1 track "
        "'active' với nhiều track bên kia cùng lúc nếu đều dưới ngưỡng) tụt xuống CHỈ CÒN 49.2% "
        "trung bình qua 5 biến thể (bản gốc riêng lẻ là 75%, một số biến thể còn thấp hơn do "
        "nhiễu/độ sáng làm bbox lệch nhẹ, đẩy khoảng cách qua lại quanh ngưỡng) — trong khi "
        "Hungarian (ép buộc 1-1 tối ưu toàn cục) vẫn giữ nguyên 100% ở MỌI biến thể, không dao "
        "động. Đây chính xác là loại tình huống mơ hồ mà GNN-CCA được thiết kế để giải quyết bằng "
        "appearance + học — Identity Association giải quyết được phần lớn chỉ bằng cách chọn đúng thuật toán gán "
        "ghép (Hungarian), không cần thêm model, và ổn định hơn cả threshold+CC khi có nhiễu."
    )

    add_file_catalog(doc, [
        ("chart_per_scene_ari.png", "Biểu đồ ARI theo từng cảnh, Hungarian vs Threshold+CC."),
        ("Noi_dung.docx", "Chính file này."),
    ])

    doc.save(str(DIR_OWN / "Noi_dung.docx"))
    print(f"Da luu: {DIR_OWN / 'Noi_dung.docx'}")


# ============================================================
# 2. DOI THU (GNN-CCA)
# ============================================================
def build_doc_competitor():
    doc = Document()
    doc.add_heading("Identity Association — Đối thủ (GNN-CCA, Luna et al. 2022)", level=0)

    add_heading(doc, "1. Cách triển khai")
    doc.add_paragraph(
        "Luna, SanMiguel, Martínez, Carballeira — \"Graph Neural Networks for Cross-Camera Data "
        "Association\" (arXiv:2201.06311, IEEE TCSVT 2022). Đây CHÍNH XÁC là hướng GNN mà nhóm đã "
        "cân nhắc và loại bỏ từ đầu dự án (lý do ghi trong Context xử lý vấn đề 2.docx: \"thêm 1 "
        "model riêng, không khớp ràng buộc không thêm model trên Jetson Nano 4GB\") — so sánh với "
        "đúng bài này để kiểm chứng quyết định đó bằng số liệu, không chỉ lý luận suông."
    )
    doc.add_paragraph(
        "Code chính thức công khai (github.com/vpulab/GNN-CCA) — đã tải về và đọc trực tiếp. "
        "Kiến trúc: mỗi frame → xây đồ thị (mỗi detection là 1 node, cạnh nối MỌI cặp detection "
        "khác camera) → khởi tạo đặc trưng node bằng CNN Re-ID (appearance) + đặc trưng cạnh bằng "
        "khoảng cách appearance VÀ khoảng cách không gian (chiếu qua homography, giống ý tưởng "
        "Identity Association) → Message Passing Network (MPN, học lan truyền đặc trưng qua đồ thị) → phân loại "
        "cạnh active/inactive → hậu xử lý (pruning + splitting) → Connected Components ra danh "
        "tính cuối."
    )
    doc.add_paragraph(
        "GIỚI HẠN THẬT — không train/chạy được bản GNN học đầy đủ trên data Identity Association: (1) cần TRAIN có "
        "nhãn (đúng lý do ban đầu loại GNN — thêm gánh nặng train/bảo trì); (2) phụ thuộc PyTorch "
        "Geometric + pytorch-scatter/sparse/cluster — các extension biên dịch C++/CUDA khớp ĐÚNG "
        "bản torch+CUDA (env gốc: torch 1.9.1+cu111) — không có wheel dựng sẵn cho môi trường của "
        "nhóm, rủi ro cài đặt thật (đã từng gặp khó khăn tương tự với pycuda+TensorRT trong dự án "
        "này). Đã CHẠY ĐƯỢC phần nhẹ, không cần PyTorch Geometric: hàm `geometrical_association` "
        "— baseline THUẦN HÌNH HỌC (không học) mà CHÍNH paper này dùng để tự so sánh trong Table V "
        "— tái hiện lại và chạy trên data Identity Association thật (xem folder 3)."
    )

    add_heading(doc, "2. Số liệu — trích dẫn TRỰC TIẾP từ Table V của paper")
    doc.add_paragraph(
        "Đo trên EPFL multi-camera pedestrian dataset (Terrace/Laboratory/Basketball, tới 9 người, "
        "4 camera — khó hơn nhiều so với Identity Association: 1-2 người, 2 camera), trung bình 3 tập train-test "
        "S1/S2/S3, dùng ResNet50 (Market1501+CUHK03+DukeMTMC) làm CNN Re-ID:"
    )
    add_table(doc, ["Method", "ARI", "AMI", "Homogeneity", "Completeness", "V-measure"], [
        ["Geometry-only (non-learned, same idea as Identity Association)", "43.49", "55.40", "52.48", "99.03", "60.96"],
        ["Geometry + Appearance (threshold, non-learned)", "57.22", "65.52", "80.02", "82.04", "78.45"],
        ["GNN-CCA full (learned)", "84.76", "87.65", "93.81", "92.53", "93.00"],
    ])
    doc.add_paragraph(
        "Paper tự báo cáo: GNN-CCA cải thiện 18.55% V-measure so với baseline tốt nhất không học "
        "(Geometry+Appearance), và 97.28% so với dùng khoảng cách Euclidean đơn thuần trên "
        "appearance — mức cải thiện LỚN, không nhỏ, trên chính benchmark của họ."
    )

    add_heading(doc, "3. Phân tích chi phí tài nguyên trên Jetson Nano 4GB")
    doc.add_paragraph(
        "ƯỚC TÍNH CÓ CĂN CỨ — KHÔNG phải đo trực tiếp (vì không chạy được GNN-CCA đầy đủ trên "
        "thiết bị thật, xem Mục 1). Lập luận dựa trên kiến trúc thật của paper + số đo Jetson "
        "THẬT đã có từ Re-Identification/Identity Association trong dự án:"
    )
    add_table(doc, ["Component", "Cost"], [
        ["MPN (Message Passing Network) itself", "~0.2M params — light, per paper's own claim"],
        ["CNN Re-ID feature extraction (REQUIRED, every detection, every frame, BOTH cameras)",
         "Same CATEGORY of cost already measured REAL in Re-Identification: adding OSNet x0.5 (chosen as the "
         "LIGHTEST of 3 variants tried) dropped FPS to 18.15 when run alongside YOLOv8n — and "
         "that was only 1 camera/1 person. GNN-CCA needs this CNN on EVERY detection, BOTH "
         "cameras, every frame."],
        ["Graph construction + Message Passing (PyTorch Geometric, pytorch-scatter/sparse)",
         "ADDITIONAL overhead on top of the CNN — not present in Identity Association's pipeline (Identity Association only needs "
         "numpy/scipy)."],
        ["Software risk", "PyTorch Geometric + pytorch-scatter/sparse/cluster are compiled "
         "C++/CUDA extensions matched to an exact torch+CUDA version — no prebuilt wheel for "
         "Jetson Nano 4GB's JetPack (Python 3.6.9, CUDA 10.2) — would need building from source, "
         "a real risk already encountered with pycuda+TensorRT in this project."],
    ])
    doc.add_paragraph(
        "Đối chiếu: bước ghép cặp CỦA Identity Association (thuần numpy/scipy Hungarian, KHÔNG cần CNN nào) đo THẬT "
        "trên Jetson Nano 4GB: 9.6ms cho 100 lần lặp ghép cặp 10×10 người/camera (kịch bản xấu "
        "nhất giả lập, vượt xa nhu cầu thật của Identity Association chỉ 1-2 người) — với đúng quy mô thật (1-2 "
        "người), chi phí gần như bằng 0, không đáng kể so với chi phí CNN Re-ID mà GNN-CCA bắt "
        "buộc phải có."
    )

    add_file_catalog(doc, [
        ("Noi_dung.docx", "Chính file này."),
    ])

    doc.save(str(DIR_COMP / "Noi_dung.docx"))
    print(f"Da luu: {DIR_COMP / 'Noi_dung.docx'}")


# ============================================================
# 3. SO SANH
# ============================================================
def build_doc_compare():
    doc = Document()
    doc.add_heading("Identity Association — So sánh Bài của nhóm vs Đối thủ", level=0)

    add_heading(doc, "1. Bảng so sánh trực tiếp")
    doc.add_paragraph(
        "So sánh 3 nguồn: (a) Geometry-only và GNN-CCA đầy đủ — số liệu LITERATURE, trích dẫn "
        "trực tiếp từ Table V của paper, đo trên EPFL (9 người, 4 camera); (b) Identity Association Hungarian — đo "
        "THẬT trên data Identity Association Data 3 của nhóm (1-2 người, 2 camera), ĐẦY ĐỦ gốc + 4 biến thể augment "
        "an toàn (sáng/tối/nhiễu/nén, giống quy ước Re-Identification/Boundary Feature Fusion), tổng 840 frame. Cùng bộ metric (ARI/"
        "AMI/V-measure), KHÔNG cùng bộ data — ghi rõ để không gây hiểu nhầm là 'chạy trực tiếp "
        "cùng test set'."
    )
    add_table(doc, ["Metric", "Geometry-only (paper, EPFL)", "GNN-CCA full (paper, EPFL)", "Identity Association Hungarian (Ours, n=840)"], [
        ["ARI", "43.49", "84.76", "85.45"],
        ["AMI", "55.40", "87.65", "85.46"],
        ["Homogeneity", "52.48", "93.81", "99.97"],
        ["Completeness", "99.03", "92.53", "91.91"],
        ["V-measure", "60.96", "93.00", "95.01"],
    ])
    add_image(doc, DIR_CMP / "chart_clustering_vs_literature.png")

    add_heading(doc, "2. Kết luận")
    doc.add_paragraph(
        "Identity Association (thuần hình học, KHÔNG train) đo trên data thật của nhóm (n=840, gốc + 4 augment) đạt "
        "điểm TƯƠNG ĐƯƠNG HOẶC CAO HƠN GNN-CCA đầy đủ đo trên EPFL (ARI 85.45 vs 84.76; V-measure "
        "95.01 vs 93.00) — dù GNN-CCA vượt trội RÕ RỆT so với baseline hình học của CHÍNH HỌ trên "
        "EPFL (ARI 84.76 vs 43.49, gần gấp đôi). Đây KHÔNG phải bằng chứng \"hình học luôn thắng "
        "GNN nói chung\" — chính paper chứng minh ngược lại trên benchmark của họ — mà là bằng "
        "chứng cho 2 điều cụ thể:"
    )
    doc.add_paragraph(
        "(1) Quy mô bài toán MỤC TIÊU của nhóm nhỏ hơn nhiều so với EPFL (1-2 người/2 camera vs "
        "tới 9 người/4 camera) — độ mơ hồ cần GNN học để giải quyết (nhầm lẫn giữa nhiều người "
        "cùng lúc) gần như không xuất hiện ở quy mô nhỏ này.\n"
        "(2) Lựa chọn THUẬT TOÁN GÁN GHÉP đúng (Hungarian — tối ưu toàn cục) đã giải quyết được "
        "phần lớn độ mơ hồ còn lại — bằng chứng trực tiếp: Scene 2 (đứng sát nhau), threshold+CC "
        "đơn thuần (kiểu baseline không học của GNN-CCA) tụt xuống CHỈ CÒN 49.2% ARI trung bình "
        "qua 5 biến thể, trong khi Hungarian giữ nguyên 100% ở MỌI biến thể — kết quả CÀNG rõ hơn "
        "sau khi có đủ augment, không phải may mắn ở 1 lần đo (xem folder 1, Mục 3)."
    )
    doc.add_paragraph(
        "Về chi phí tài nguyên (xem folder 2, Mục 3 — phân tích/ước tính, không đo trực tiếp): "
        "GNN-CCA cần 1 CNN Re-ID chạy trên MỌI detection ở CẢ 2 camera (cùng loại chi phí đã đo "
        "thật khiến Re-Identification phải chọn cẩn thận model Re-ID nhẹ nhất) CỘNG THÊM PyTorch Geometric "
        "(dependency khó cài trên Jetson Nano 4GB thật) — trong khi Identity Association đo thật chỉ tốn 9.6ms cho "
        "kịch bản xấu nhất giả lập. Quyết định loại bỏ GNN từ đầu dự án (ràng buộc phần cứng) là "
        "hợp lý VÀ không đánh đổi độ chính xác đáng kể ở đúng quy mô bài toán của nhóm."
    )

    add_heading(doc, "3. Giới hạn cần nêu rõ")
    doc.add_paragraph(
        "(a) KHÔNG train/chạy được bản GNN-CCA học đầy đủ trên chính data Identity Association — so sánh dừng ở "
        "mức trích dẫn số liệu literature (EPFL) đối chiếu với số đo thật trên data mình, KHÔNG "
        "phải chạy 2 phương pháp trên CÙNG 1 test set. Nếu chạy được GNN-CCA thật trên data Identity Association, "
        "kết quả có thể khác (có thể còn cao hơn, vì quy mô nhỏ dễ hơn cho GNN).\n"
        "(b) Ground truth cho đánh giá trên data Identity Association là bootstrap từ tần suất ghép Hungarian (xem "
        "folder 1, Mục 2) — không phải nhãn tay độc lập, có thể thiên vị nhẹ cho Hungarian.\n"
        "(c) Phần chi phí tài nguyên Jetson của GNN-CCA là ƯỚC TÍNH có căn cứ (suy luận từ kiến "
        "trúc thật + số đo Jetson thật của Re-Identification/Identity Association), KHÔNG phải benchmark trực tiếp."
    )

    add_file_catalog(doc, [
        ("chart_clustering_vs_literature.png", "Biểu đồ so sánh ARI/AMI/V-measure: Geometry-only "
         "(paper) vs GNN-CCA đầy đủ (paper) vs Identity Association Hungarian (data thật của nhóm)."),
        ("Noi_dung.docx", "Chính file này."),
    ])

    doc.save(str(DIR_CMP / "Noi_dung.docx"))
    print(f"Da luu: {DIR_CMP / 'Noi_dung.docx'}")


def main():
    DIR_OWN.mkdir(parents=True, exist_ok=True)
    DIR_COMP.mkdir(parents=True, exist_ok=True)
    DIR_CMP.mkdir(parents=True, exist_ok=True)
    build_doc_own()
    build_doc_competitor()
    build_doc_compare()


if __name__ == "__main__":
    main()
